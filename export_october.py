#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""10/1〜10/31 の出荷日データを全ソースから抽出してデスクトップに保存する。

出力シート:
  1. 10月全件         … 全ソース・全シートの合計
  2. 出庫案件まとめ   … 宇美工場出荷分 ／ 外部倉庫出庫分 を区別して一覧

使い方: py -3 export_october.py
"""

from __future__ import annotations

import json
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from driver_sync import (  # noqa: E402
    download_share_file,
    extract_rows_from_workbook,
    SOURCE_ORDER,
    SOURCE_DISPLAY_NAMES,
)

CONFIG_PATH = Path(__file__).resolve().parent / "driver_sync_config.json"
DESKTOP = Path.home() / "Desktop"
BACKUP_FOLDER = DESKTOP / "ドライバー情報バックアップ"

EXPORT_START = date(2026, 10, 1)
EXPORT_END   = date(2026, 10, 31)

COLUMNS = ["出荷日", "着日", "案件No", "案件名", "型式", "車型",
           "会社名", "乗務員", "車番", "携帯番号", "備考", "依頼先"]
COL_WIDTHS = [10, 10, 14, 28, 20, 12, 18, 12, 12, 16, 20, 12]

# ソースごとの行色（10月全件シート用）
SOURCE_COLORS = {
    "matsuzaki": "D9EAD3",
    "nakadori":  "CFE2F3",
    "fukuoka":   "FFF2CC",
    "maruun":    "FCE5CD",
    "ks":        "EAD1DC",
}

# 宇美工場 / 外部倉庫 のセクション色
COLOR_UMI    = "D6E4F0"   # 淡青：宇美工場出荷分
COLOR_GABU   = "FEF3CD"   # 淡橙：外部倉庫出庫分
COLOR_HEADER_UMI  = "2E75B6"
COLOR_HEADER_GABU = "C55A11"


def _fetch_rows(wb, key, *, exclude_keywords=()):
    return extract_rows_from_workbook(
        wb, key, 4, EXPORT_START,
        arr_days_back=0,
        arr_window_end_override=EXPORT_END,
        window_date_field="ship",
        require_driver_info=False,
        excluded_sheet_keywords=exclude_keywords,
    )


def write_header(ws, row_idx: int):
    ws.append(COLUMNS)
    for cell in ws[row_idx]:
        cell.font  = Font(bold=True, color="FFFFFF")
        cell.fill  = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_idx].height = 20


def write_data_row(ws, values: list, fill_color: str):
    ws.append(values)
    fill = PatternFill("solid", fgColor=fill_color)
    for cell in ws[ws.max_row]:
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_section_header(ws, label: str, bg_color: str):
    """セクション区切り行（結合セル風の見出し）を挿入する。"""
    ws.append([label] + [""] * (len(COLUMNS) - 1))
    row_idx = ws.max_row
    ws.merge_cells(
        start_row=row_idx, start_column=1,
        end_row=row_idx, end_column=len(COLUMNS)
    )
    cell = ws.cell(row=row_idx, column=1)
    cell.font  = Font(bold=True, size=12, color="FFFFFF")
    cell.fill  = PatternFill("solid", fgColor=bg_color)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row_idx].height = 22


def set_col_widths(ws):
    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def row_values(row: dict) -> list:
    key = row.get("依頼先_key", "")
    return [
        row.get("出荷日", ""), row.get("着日", ""),
        row.get("案件No", ""), row.get("案件名", ""),
        row.get("型式", ""),  row.get("車型", ""),
        row.get("会社名", ""), row.get("乗務員", ""),
        row.get("車番", ""),  row.get("携帯番号", ""),
        row.get("備考", ""),  SOURCE_DISPLAY_NAMES.get(key, key),
    ]


def main() -> int:
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    sources = {s["key"]: s["share_url"] for s in config.get("sources", [])}

    umi_rows:  list[dict] = []   # 宇美工場（車両依頼書ベース）
    gabu_rows: list[dict] = []   # 外部倉庫（管理シートベース）

    for key in SOURCE_ORDER:
        share_url = sources.get(key)
        if not share_url:
            print(f"[{key}] share_url なし — スキップ")
            continue
        print(f"[{key}] ダウンロード中...")
        try:
            content = download_share_file(share_url)
        except Exception as e:
            print(f"[{key}] 取得失敗: {e}")
            continue

        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)

        # 宇美工場：「管理」の付くシートを除外 → 車両依頼書のみ
        u_rows = _fetch_rows(wb, key, exclude_keywords=("管理",))
        for r in u_rows:
            r["依頼先_key"] = key
        umi_rows.extend(u_rows)
        print(f"  宇美工場: {len(u_rows)} 件")

        # 外部倉庫：「依頼書」の付くシートを除外 → 管理シート系のみ
        g_rows = _fetch_rows(wb, key, exclude_keywords=("依頼書",))
        for r in g_rows:
            r["依頼先_key"] = key
        gabu_rows.extend(g_rows)
        print(f"  外部倉庫: {len(g_rows)} 件")

        wb.close()

    all_rows = sorted(umi_rows + gabu_rows,
                      key=lambda r: (r.get("出荷日") or "", r.get("依頼先_key") or ""))

    # ── Excel 出力 ─────────────────────────────────────────────
    out_wb = openpyxl.Workbook()

    # ── Sheet1: 10月全件 ────────────────────────────────────────
    ws_all = out_wb.active
    ws_all.title = "10月全件"
    write_header(ws_all, 1)
    for row in all_rows:
        key = row.get("依頼先_key", "")
        write_data_row(ws_all, row_values(row), SOURCE_COLORS.get(key, "FFFFFF"))
    set_col_widths(ws_all)
    print(f"\n[10月全件] {len(all_rows)} 件")

    # ── Sheet2: 出庫案件まとめ ───────────────────────────────────
    ws_sum = out_wb.create_sheet("出庫案件まとめ")

    # 宇美工場 セクション
    write_section_header(ws_sum, f"■ 宇美工場 出荷分  （{len(umi_rows)} 件）", COLOR_HEADER_UMI)
    write_header(ws_sum, ws_sum.max_row + 1)
    umi_sorted = sorted(umi_rows, key=lambda r: (r.get("出荷日") or "", r.get("依頼先_key") or ""))
    for row in umi_sorted:
        write_data_row(ws_sum, row_values(row), COLOR_UMI)

    # 区切り（空行）
    ws_sum.append([""] * len(COLUMNS))
    ws_sum.row_dimensions[ws_sum.max_row].height = 10

    # 外部倉庫 セクション
    write_section_header(ws_sum, f"■ 外部倉庫 出庫分  （{len(gabu_rows)} 件）", COLOR_HEADER_GABU)
    write_header(ws_sum, ws_sum.max_row + 1)
    gabu_sorted = sorted(gabu_rows, key=lambda r: (r.get("出荷日") or "", r.get("依頼先_key") or ""))
    for row in gabu_sorted:
        write_data_row(ws_sum, row_values(row), COLOR_GABU)

    set_col_widths(ws_sum)
    print(f"[出庫案件まとめ] 宇美工場 {len(umi_rows)} 件 ／ 外部倉庫 {len(gabu_rows)} 件")

    # ── 保存 ────────────────────────────────────────────────────
    BACKUP_FOLDER.mkdir(parents=True, exist_ok=True)
    out_path = BACKUP_FOLDER / "ドライバー情報_2026-10_全件.xlsx"
    out_wb.save(out_path)
    print(f"\n保存完了: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
